from copy import deepcopy

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference


async def create_bar(sheet_name="bar.xlsx", max_col=6, *, rows):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    for row in rows:
        ws.append(row)
    chart1 = BarChart()
    chart1.type = "col"
    chart1.style = 10
    chart1.title = "Week Tracking"
    chart1.y_axis.title = 'Devoted time'
    chart1.x_axis.title = 'Days of the week'

    data = Reference(ws, min_col=2, min_row=1, max_row=8, max_col=max_col)
    cats = Reference(ws, min_col=1, min_row=2, max_row=8)
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    chart1.shape = 4
    ws.add_chart(chart1, "A10")
    await horizontal_bar_chart(chart1, ws)
    await stacked_chart(chart1, ws)
    await percent_stacked_hart(chart1, ws)
    wb.save(sheet_name)


async def horizontal_bar_chart(chart1, ws):
    chart2 = deepcopy(chart1)
    chart2.style = 10
    chart2.type = "bar"
    chart2.title = "Horizontal Bar Chart"
    ws.add_chart(chart2, "I10")


#
async def stacked_chart(chart1, ws):
    chart3 = deepcopy(chart1)
    chart3.type = "col"
    chart3.style = 10
    chart3.grouping = "stacked"
    chart3.overlap = 100
    chart3.title = 'Stacked Chart'
    ws.add_chart(chart3, "A27")


async def percent_stacked_hart(chart1, ws):
    chart4 = deepcopy(chart1)
    chart4.type = "bar"
    chart4.style = 10
    chart4.grouping = "percentStacked"
    chart4.overlap = 100
    chart4.title = 'Percent Stacked Chart'
    ws.add_chart(chart4, "I27")